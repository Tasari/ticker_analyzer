"""Bounded XLSX loading and statement cell parsing."""

from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime
from io import BytesIO
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ticker_analyzer.portfolio.statement_models import AccountStatementError

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 75 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 250


def validate_xlsx_payload(payload: bytes) -> None:
    if not payload:
        raise AccountStatementError("The uploaded file is empty.")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise AccountStatementError("The workbook is larger than the 10 MB upload limit.")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise AccountStatementError("The workbook contains too many internal files.")
            if sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES:
                raise AccountStatementError("The expanded workbook is too large to process safely.")
            if "xl/workbook.xml" not in archive.namelist():
                raise AccountStatementError("The uploaded file is not an XLSX workbook.")
    except BadZipFile as exc:
        raise AccountStatementError("The uploaded file is not an XLSX workbook.") from exc


def _load_statement_workbook(payload: bytes) -> Any:
    validate_xlsx_payload(payload)
    try:
        return load_workbook(
            BytesIO(payload),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise AccountStatementError("The uploaded file is not a readable XLSX workbook.") from exc


def _key_value_rows(worksheet: Any) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for row in worksheet.iter_rows(min_col=1, max_col=2, values_only=True):
        key = _optional_text(row[0])
        if key:
            values[key] = row[1]
    return values


def _row_value(row: tuple[Any, ...], index: int) -> Any:
    return row[index] if index < len(row) else None


def _required_datetime(summary: dict[str, Any], key: str) -> datetime:
    value = _parse_statement_datetime(summary.get(key))
    if value is None:
        raise AccountStatementError(f"Portfolio analysis is unavailable: missing {key}.")
    return value


def _required_number(summary: dict[str, Any], key: str) -> float:
    if summary.get(key) is None:
        raise AccountStatementError(f"Portfolio analysis is unavailable: missing {key}.")
    return _number(summary[key])


def _number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = _optional_text(value)
    if not text or text == "-":
        return 0.0
    negative = text.startswith("(") and text.endswith(")")
    normalized = text.strip("()").replace(",", "").replace("$", "")
    try:
        parsed = float(normalized)
    except ValueError:
        return 0.0
    return -parsed if negative else parsed


def _worksheet_shape(worksheet: Any) -> tuple[int, int]:
    if worksheet.max_row and worksheet.max_column:
        return max(worksheet.max_row - 1, 0), worksheet.max_column
    row_count = 0
    column_count = 0
    for row in worksheet.iter_rows(values_only=True):
        row_count += 1
        column_count = max(column_count, len(row))
    return max(row_count - 1, 0), column_count


def _unique_headers(values: tuple[Any, ...]) -> tuple[str, ...]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = _optional_text(value) or f"Column {index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base} ({counts[base]})")
    return tuple(headers)


def _optional_text(value: Any) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None


def _parse_statement_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = _optional_text(value)
    if not text:
        return None
    for pattern in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    return None


def _sheet_rows(worksheet: Any) -> tuple[dict[str | None, int], Iterator[tuple[Any, ...]]]:
    rows = worksheet.iter_rows(values_only=True)
    columns = {_optional_text(value): index for index, value in enumerate(next(rows, ()))}
    return columns, rows
